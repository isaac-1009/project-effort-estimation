#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
工时估算Excel读取脚本

用法：
    python scripts/read_excel.py <excel文件路径>
    python scripts/read_excel.py data/current/estimate.xlsx

重构要点：
- 拆分为 load_worksheet / extract_rows / group_by_system / format_report，
  read_excel 现在返回结构化结果（dict），便于被其他脚本复用，而非纯 print。
- 修复了 "其他" 工时行序号显示为 None 的瑕疵（合并单元格非首行序号为空）。
- 描述截断长度等 magic number 提取为常量。
"""

from __future__ import annotations

import argparse
import logging
import os
import sys
from typing import Any

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from logging_setup import setup_logging  # noqa: E402

setup_logging()
log = logging.getLogger(__name__)

SHEET_NAME = "工时估算"
DESC_TRUNCATE = 60  # 描述显示截断长度（字符）

# 列布局（须与 generate_excel.py 的 8 列结构保持一致）
# A=#, B=模块, C=功能, D=需求类型, E=描述, F=系统, G=工时, H=备注
COL_SEQ = 0
COL_MOD = 1
COL_NAME = 2
COL_TYPE = 3
COL_DESC = 4
COL_SYS = 5
COL_EFFORT = 6
COL_NOTE = 7


def load_worksheet(file_path: str):
    """
    加载 Excel 的工时估算工作表。

    Raises:
        FileNotFoundError: 文件不存在
        KeyError: 工作表 "工时估算" 不存在
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")
    wb = openpyxl.load_workbook(file_path)
    if SHEET_NAME not in wb.sheetnames:
        raise KeyError(f"工作表 '{SHEET_NAME}' 不存在，可用工作表: {wb.sheetnames}")
    return wb, wb[SHEET_NAME]


def _is_total_row(effort_cell: Any) -> bool:
    """识别总计行：工时单元格是公式字符串（含 '='）。"""
    return isinstance(effort_cell, str) and "=" in effort_cell


def extract_rows(ws) -> list[dict]:
    """
    从工作表抽取数据行（跳过表头与总计行，去重合并单元格产生的重复行）。

    合并单元格的非首行在各列会是 None；本函数用"前向填充"还原系统/序号，
    避免下游显示 None。
    """
    rows: list[dict] = []
    seen_module_effort: set[tuple] = set()  # 去重：(module, effort)
    last_seq = None
    last_system = None
    last_module = None
    last_name = None
    last_type = None

    for raw in ws.iter_rows(min_row=2, values_only=True):
        # 安全取列（防御列数不足）
        seq = raw[COL_SEQ] if len(raw) > COL_SEQ else None
        module = raw[COL_MOD] if len(raw) > COL_MOD else None
        name = raw[COL_NAME] if len(raw) > COL_NAME else None
        req_type = raw[COL_TYPE] if len(raw) > COL_TYPE else None
        desc = raw[COL_DESC] if len(raw) > COL_DESC else None
        system = raw[COL_SYS] if len(raw) > COL_SYS else None
        effort = raw[COL_EFFORT] if len(raw) > COL_EFFORT else None
        note = raw[COL_NOTE] if len(raw) > COL_NOTE else None

        if effort is None:
            continue
        if _is_total_row(effort):
            continue

        # 前向填充：合并单元格非首行的 #/模块/功能/需求类型/系统 为 None，沿用上一个有效值
        # 注意：desc 不前向填充（合并的描述在首行已展示，非首行留空避免重复）
        if seq is not None:
            last_seq = seq
        if module is not None and str(module).strip():
            last_module = module
        if name is not None and str(name).strip():
            last_name = name
        if req_type is not None and str(req_type).strip():
            last_type = req_type
        if system is not None and str(system).strip() and system != "-":
            last_system = system

        seq_val = last_seq
        module_str = last_module or ""
        name_str = last_name or ""
        type_str = last_type or ""
        # 其他工时行的 system 为 "-"，保持 "-" 以便 group_by_system 归入"其他"
        if system == "-":
            system_val = "-"
        else:
            system_val = last_system or ""
        desc_str = desc if desc else ""  # 不前向填充
        note_str = note if note else ""

        # 去重：功能名为空（"其他"工时的后续行）且 (module, effort) 已见过则跳过
        if not str(name_str).strip() and module_str:
            if isinstance(effort, (int, float)):
                key = (module_str, effort)
                if key in seen_module_effort:
                    continue
                seen_module_effort.add(key)

        rows.append({
            "seq": seq_val,
            "system": system_val or "",
            "module": module_str,
            "name": name_str,
            "type": type_str,
            "effort": effort,
            "desc": desc_str,
            "note": note_str,
        })

    return rows


def group_by_system(rows: list[dict]) -> dict[str, list[dict]]:
    """
    按 system 分组（保留首次出现顺序），返回 {system: [items]}。
    新结构下其他工时行的 system 为 "-"，统一归入 "其他" 分组。
    """
    systems: dict[str, list[dict]] = {}
    for r in rows:
        # "-" 或空字符串都视为"其他"（其他工时行）
        sys_key = r["system"] if r["system"] and r["system"] != "-" else "其他"
        systems.setdefault(sys_key, []).append(r)
    return systems


def _truncate(text: str, limit: int = DESC_TRUNCATE) -> str:
    return text if len(text) <= limit else text[:limit] + "..."


def _format_system_block(system: str, items: list[dict]) -> tuple[str, float]:
    """格式化单个系统的明细，返回 (文本块, 系统小计)。"""
    lines = [f"\n【{system}】"]
    system_total = 0.0
    for item in items:
        effort = item["effort"]
        if not isinstance(effort, (int, float)):
            continue
        system_total += effort
        module_str = f" - {item['module']}" if item["module"] and item["module"] != "-" else ""
        name_str = item["name"] if item["name"] and item["name"] != "-" else ""
        seq_str = f"{item['seq']}." if item["seq"] is not None else ""
        lines.append(f"  {seq_str} {name_str}{module_str} | {effort}人天".strip())
        if item["desc"]:
            lines.append(f"     {_truncate(item['desc'])}")
        if item.get("note"):
            lines.append(f"     备注: {_truncate(item['note'])}")
    if system_total > 0:
        lines.append(f"  小计: {system_total}人天")
    return "\n".join(lines), system_total


def format_report(file_path: str, rows: list[dict],
                  systems: dict[str, list[dict]], total_effort: float) -> str:
    """生成完整的人类可读报告文本。"""
    lines = ["=" * 80, f"工时估算文件: {file_path}", "=" * 80]

    # 普通系统在前，"其他" 最后
    ordered = [s for s in systems if s != "其他"]
    if "其他" in systems:
        ordered.append("其他")

    for system in ordered:
        block, _ = _format_system_block(system, systems[system])
        lines.append(block)

    lines.append("\n" + "=" * 80)
    lines.append(f"总计工时: {total_effort}人天")
    lines.append(f"功能数量: {len(rows)}项")
    lines.append("=" * 80)
    return "\n".join(lines)


def read_excel(file_path: str, *, print_report: bool = True) -> dict:
    """
    读取工时估算 Excel 文件。

    Args:
        file_path:    Excel 文件路径
        print_report: 是否打印人类可读报告（默认 True，向后兼容）

    Returns:
        {"rows": [...], "systems": {...}, "total_effort": float}

    Raises:
        FileNotFoundError, KeyError
    """
    wb, ws = load_worksheet(file_path)
    try:
        rows = extract_rows(ws)
        if not rows:
            log.warning("未找到数据行")
            return {"rows": [], "systems": {}, "total_effort": 0}

        systems = group_by_system(rows)
        total_effort = sum(
            r["effort"] for r in rows if isinstance(r["effort"], (int, float))
        )

        if print_report:
            print(format_report(file_path, rows, systems, total_effort))

        return {"rows": rows, "systems": systems, "total_effort": total_effort}
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="读取工时估算 Excel 文件")
    parser.add_argument("file_path", help="Excel 文件路径")
    args = parser.parse_args()

    try:
        read_excel(args.file_path)
        return 0
    except FileNotFoundError as e:
        log.error("%s", e)
        return 1
    except KeyError as e:
        log.error("%s", e)
        return 1
    except Exception:
        log.exception("读取文件失败")
        return 1


if __name__ == "__main__":
    sys.exit(main())
