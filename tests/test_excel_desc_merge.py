"""generate_excel 描述列（E）合并逻辑的测试。

需求：同一功能跨多行（企微/PC后台）时，若各行描述文本相同，
E 列应纵向合并；描述不同则保持每行独立。
"""

import json

import openpyxl
import pytest

from generate_excel import desc_merge_value, generate_excel
from read_excel import read_excel


class TestDescMergeValue:
    def test_all_same_returns_value(self):
        assert desc_merge_value(["同一段描述", "同一段描述"]) == "同一段描述"

    def test_single_row_returns_value(self):
        assert desc_merge_value(["只此一行"]) == "只此一行"

    def test_mixed_returns_none(self):
        assert desc_merge_value(["A描述", "B描述"]) is None

    def test_empty_string_returns_none(self):
        assert desc_merge_value(["", ""]) is None

    def test_empty_list_returns_none(self):
        assert desc_merge_value([]) is None


@pytest.fixture
def merge_metadata_path(tmp_path):
    """两功能四行：功能A 两行描述相同（应合并），功能B 两行描述不同（不合并）。"""
    metadata = {
        "project_id": "合并测试_20260902",
        "name": "合并测试",
        "created_date": "2026-09-02",
        "total_effort": 7,
        "description": "测试描述列合并",
        "systems": ["企微", "PC后台"],
        "tags": [],
        "tech_stack": {
            "frontend": "H5",
            "backend": "Java",
            "database": "MySQL",
        },
        "features": [
            {
                "module": "功能开发",
                "name": "功能A",
                "requirement_type": "功能优化",
                "description": "相同的描述文本",
                "sub_systems": [
                    {"system": "企微", "effort_days": 2},
                    {"system": "PC后台", "effort_days": 3},
                ],
            },
            {
                "module": "功能开发",
                "name": "功能B",
                "requirement_type": "功能优化",
                "description": "各自的描述",
                "sub_systems": [
                    {"system": "企微", "effort_days": 1, "description": "企微侧描述"},
                    {"system": "PC后台", "effort_days": 1, "description": "后台侧描述"},
                ],
            },
        ],
        "summary": {
            "user_development": 3,
            "admin_development": 4,
            "total_development_effort": 7,
            "testing_effort": 0,
            "other_fixed": {
                "testing": False,
                "acceptance": 0,
                "communication": 0,
                "ui_ux_design": 0,
            },
            "total_effort": 7,
        },
    }
    path = tmp_path / "metadata.json"
    path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
    return path


class TestExcelDescMerge:
    def test_same_desc_merged_and_diff_desc_not(self, tmp_path, merge_metadata_path):
        out = generate_excel(str(merge_metadata_path), str(tmp_path / "out.xlsx"))
        wb = openpyxl.load_workbook(out)
        ws = wb["工时估算"]
        merged = {str(r) for r in ws.merged_cells.ranges}

        # 功能A（Excel 第 2-3 行）：描述相同 -> E 列合并（随 A/C 一起并入合并组）
        assert "E2:E3" in merged
        # 功能B（第 4-5 行）：描述不同 -> E 列不合并
        assert "E4:E5" not in merged
        # 描述首行保留文本
        assert ws.cell(row=2, column=5).value == "相同的描述文本"
        assert ws.cell(row=4, column=5).value == "企微侧描述"

    def test_totals_unchanged_by_merging(self, tmp_path, merge_metadata_path):
        out = generate_excel(str(merge_metadata_path), str(tmp_path / "out.xlsx"))
        result = read_excel(out, print_report=False)
        # 合并只影响展示，工时合计不受影响：2+3+1+1
        assert result["total_effort"] == 7.0
        assert len(result["rows"]) == 4
