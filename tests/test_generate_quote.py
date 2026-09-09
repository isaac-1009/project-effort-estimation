"""generate_quote 报价版生成器的测试。

核心约束：给定目标总工时，逐行上浮（0.5 粒度、不低于基准）后，
Excel 实算合计必须恰好等于目标值。
"""

import json
import math
import os

import openpyxl
import pytest

from generate_quote import (
    base_totals,
    distribute_uplift,
    excel_roundup,
    generate_quote,
    solve_dev_target,
)
from metadata_schema import ValidationError
from read_excel import read_excel

FACILITY_METADATA = os.path.join(
    "data", "estimations", "Facility 三期需求_20260902", "metadata.json"
)


def _metadata(dev_values, *, testing=True):
    """按开发行工时列表构造最小 metadata（全部为企微单行功能，验收5+沟通8）。"""
    features = [
        {
            "module": "功能开发",
            "name": f"功能{i}",
            "requirement_type": "新功能",
            "description": "d",
            "complexity": "低",
            "system": "企微",
            "effort_days": v,
        }
        for i, v in enumerate(dev_values)
    ]
    dev = sum(dev_values)
    testing_effort = math.ceil(dev * 0.15) if testing else 0
    other_fixed = {"acceptance": 5, "communication": 8, "ui_ux_design": 0}
    if not testing:
        other_fixed["testing"] = False
    total = dev + testing_effort + 13
    return {
        "project_id": "报价测试_20260902",
        "name": "报价测试",
        "created_date": "2026-09-02",
        "total_effort": total,
        "description": "测试",
        "systems": ["企微"],
        "tags": [],
        "tech_stack": {"frontend": "H5", "backend": "Java", "database": "MySQL"},
        "features": features,
        "summary": {
            "user_development": dev,
            "admin_development": 0,
            "total_development_effort": dev,
            "testing_effort": testing_effort,
            "other_fixed": other_fixed,
            "total_effort": total,
        },
    }


class TestSolveDevTarget:
    def test_facility_case(self):
        # Facility 三期：开发 153 + 测试 23 + 固定 13 = 189，目标 249.5
        dev = solve_dev_target(153, 13, True, 249.5)
        assert dev == 205.5
        assert 205.5 + excel_roundup(205.5 * 0.15) + 13 == 249.5

    def test_no_testing_direct_subtract(self):
        assert solve_dev_target(100, 13, False, 200) == 187.0

    def test_unreachable_target_raises(self):
        # dev 与 ROUNDUP 联动存在 0.5 空隙：need=236.5 可达(d=205.5)、237 可达(d=206)，
        # 238 不可达（207 直接到 239）
        with pytest.raises(ValidationError, match="不可达"):
            solve_dev_target(153, 13, True, 238 + 13)


class TestDistributeUplift:
    def test_exact_hit_and_grid(self):
        rows = distribute_uplift([5, 2, 2, 1, 10], 27.0)
        assert sum(rows) == 27.0
        assert all(v * 2 == int(v * 2) for v in rows), "必须落在 0.5 网格"

    def test_no_row_below_base(self):
        base = [3, 1, 1, 8, 2]
        rows = distribute_uplift(base, sum(base) * 1.4)
        assert all(v >= b for v, b in zip(rows, base))

    def test_below_base_rejected(self):
        with pytest.raises(ValidationError, match="只支持上浮"):
            distribute_uplift([5, 5], 8.0)

    def test_large_uplift_diff_exceeds_rows(self):
        # 差额单元数超过行数时循环补足
        rows = distribute_uplift([1.0, 1.0], 6.0)
        assert sum(rows) == 6.0
        assert all(v >= b for v, b in zip(rows, [1.0, 1.0]))


def _quote_excel_dev_sum(out_path) -> float:
    """读取报价版 Excel，模拟公式求实算合计（开发 + ROUNDUP(开发×15%) + 数值固定行）。"""
    ws = openpyxl.load_workbook(str(out_path))["工时估算"]
    numeric = [
        ws.cell(row=r, column=7).value
        for r in range(2, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=7).value, (int, float))
    ]
    dev = sum(numeric[:-2])  # 末两行数值 = 验收 + 沟通
    fixed = sum(numeric[-2:])
    return dev + math.ceil(dev * 0.15) + fixed


class TestGenerateQuoteIntegration:
    def _write(self, tmp_path, metadata):
        path = tmp_path / "metadata.json"
        path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")
        return path

    def test_end_to_end_formula_total_equals_target(self, tmp_path):
        metadata = _metadata([3, 5, 2, 4, 1, 8])  # 开发 23，基准合计 40
        target = 56.5
        out = tmp_path / "quote.xlsx"
        result = generate_quote(self._write(tmp_path, metadata), target, str(out))

        assert result["target_total"] == target
        assert result["base_total"] == 40
        assert _quote_excel_dev_sum(out) == target
        assert all(q >= b for b, q in result["rows"])
        assert all(q * 2 == int(q * 2) for _, q in result["rows"])

    def test_read_excel_loads_quote(self, tmp_path):
        # 基准：开发 9 + 测试 2 + 固定 13 = 24；目标 30 → 开发' = 14（测试计提 3）
        out = tmp_path / "quote.xlsx"
        generate_quote(self._write(tmp_path, _metadata([2, 3, 4])), 30.0, str(out))
        result = read_excel(str(out), print_report=False)
        # read_excel 数值合计 = 开发' + 固定 13（测试为公式行被跳过）
        assert result["total_effort"] == 27.0

    def test_target_below_base_raises(self, tmp_path):
        with pytest.raises(ValidationError, match="只支持上浮"):
            generate_quote(
                self._write(tmp_path, _metadata([3, 5])), 10.0, str(tmp_path / "q.xlsx")
            )

    @pytest.mark.skipif(not os.path.exists(FACILITY_METADATA), reason="归档数据不存在")
    def test_facility_real_archive_189_to_249_5(self, tmp_path):
        out = tmp_path / "quote.xlsx"
        result = generate_quote(FACILITY_METADATA, 249.5, str(out))
        assert result["base_total"] == 189
        assert result["target_total"] == 249.5
        assert result["dev_target"] == 205.5
        assert _quote_excel_dev_sum(out) == 249.5
